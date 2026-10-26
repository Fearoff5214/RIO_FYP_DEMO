"""Knowledge base ingestion (offline plane) and hybrid retrieval (Stage 4).

Ingest: parse -> chunk -> tag -> embed -> index. Retrieval: dense cosine + keyword overlap,
then a rerank that boosts exact defect/material matches. Vectors live in SQLite as JSON;
for large corpora swap the scan in `search` for a pgvector query.
"""
import math
import re
from pathlib import Path

from . import config, db, llm

DEFECT_TERMS = ["crazing", "inclusion", "patches", "pitted_surface", "rolled-in_scale", "scratches",
                "dent", "crack", "deformation"]


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return "\n".join(" | ".join(str(c) for c in row if c is not None)
                         for ws in wb for row in ws.iter_rows(values_only=True))
    raise ValueError(f"unsupported file type: {suffix}")


def chunk(text: str, max_chars: int = 700) -> list[str]:
    """One chunk per paragraph (so a chunk answers about one defect); tiny paragraphs and
    headings are folded into the next one, and over-long paragraphs are split on sentences."""
    out, carry = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) < 80:
            carry = f"{carry}\n{para}".strip()
            continue
        para = f"{carry}\n{para}".strip()
        carry = ""
        while len(para) > max_chars:
            cut = para.rfind(". ", 0, max_chars) + 1 or max_chars
            out.append(para[:cut].strip())
            para = para[cut:].strip()
        out.append(para)
    if carry:
        out.append(carry)
    return out


def auto_tags(text: str) -> dict:
    low = text.lower()
    tags = {"defects": [d for d in DEFECT_TERMS if d.replace("_", " ") in low or d in low]}
    for key in ("material", "process"):
        m = re.search(rf"{key}\s*[:=]\s*([\w\- ]+)", text, re.I)
        if m:
            tags[key] = m.group(1).strip().lower()
    return tags


def ingest_text(source: str, text: str, tags: dict | None = None) -> int:
    n = 0
    for piece in chunk(text):
        t = {**auto_tags(piece), **(tags or {})}
        db.add_chunk(source, piece, t, llm.embed(piece))
        n += 1
    return n


def ingest_path(path: Path) -> int:
    return ingest_text(path.name, parse_file(path))


def ingest_seed_dir(directory: Path | None = None) -> int:
    directory = directory or config.KB_SEED_DIR
    return sum(ingest_path(p) for p in sorted(directory.glob("*")) if p.suffix.lower() in (".txt", ".md", ".docx", ".xlsx"))


def _cos(a: list[float], b: list[float]) -> float:
    if len(a) != len(b):
        return 0.0
    na = math.sqrt(sum(x * x for x in a)) or 1.0
    nb = math.sqrt(sum(x * x for x in b)) or 1.0
    return sum(x * y for x, y in zip(a, b)) / (na * nb)


def search(query: str, defect_types: list[str] | None = None, material: str | None = None,
           min_sim: float | None = None, top_k: int | None = None, filter_tags: bool = True) -> list[dict]:
    """Hybrid search. filter_tags=False is the L4 'relax filters' step."""
    min_sim = config.RETRIEVE_MIN_SIM if min_sim is None else min_sim
    top_k = top_k or config.RETRIEVE_TOP_K
    qvec, qtok = llm.embed(query), set(llm.tokenize(query))
    hits = []
    for ch in db.all_chunks():
        if filter_tags and defect_types:
            chunk_defects = set(ch["tags"].get("defects", []))
            if chunk_defects and not chunk_defects & set(defect_types):
                continue
        dense = _cos(qvec, ch["embedding"])
        ctok = set(llm.tokenize(ch["text"]))
        keyword = len(qtok & ctok) / (len(qtok) or 1)
        score = 0.6 * dense + 0.4 * keyword
        # rerank boosts: exact defect type and material agreement
        if defect_types and set(ch["tags"].get("defects", [])) & set(defect_types):
            score += 0.15
        if material and material.lower() in ch["text"].lower():
            score += 0.05
        if score >= min_sim:
            hits.append({"id": ch["id"], "source": ch["source"], "text": ch["text"],
                         "score": round(score, 4), "tags": ch["tags"]})
    hits.sort(key=lambda h: h["score"], reverse=True)
    return hits[:top_k]
