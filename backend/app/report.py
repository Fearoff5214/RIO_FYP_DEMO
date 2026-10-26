"""Stage 7: compose report data and render PDF, Excel and Word files."""
from pathlib import Path

from . import config

REQUIRED_KEYS = ("job_id", "part", "result", "defects", "dimensional")


def compose(state: dict) -> dict:
    ok = not state.get("defects")
    sel = state.get("selected_cause")
    return {
        "job_id": state["job_id"],
        "part": state.get("part") or {"part_number": state.get("part_number")},
        "result": "OK" if ok else "NG",
        "defects": state.get("defects", []),
        "dimensional": state.get("dimensional") or {"status": "NOT_CHECKED"},
        "flags": state.get("flags", []),
        "rca": None if ok else {"candidates": state.get("rca", []), "selected": sel},
        "approver_comments": state.get("approver_comments", []),
        "revision": state.get("report_revision", 1),
    }


def missing(data: dict) -> list[str]:
    return [k for k in REQUIRED_KEYS if data.get(k) in (None, "")]


def _lines(data: dict) -> list[tuple[str, str]]:
    part = data["part"]
    rows = [("Job", data["job_id"]), ("Part", f"{part.get('part_number', '?')} {part.get('name', '')}".strip()),
            ("Result", data["result"]), ("Dimensional", data["dimensional"].get("status", "?"))]
    for d in data["defects"]:
        rows.append(("Defect", f"{d['cls']}  conf {d['conf']:.0%}  severity {d.get('severity', '-')}"))
    sel = (data.get("rca") or {}).get("selected")
    if sel:
        rows += [("Root cause", sel["cause"]), ("Effect", sel.get("effect", "")),
                 ("Precautions", "; ".join(sel.get("precautions", []))),
                 ("Chosen by", f"{sel.get('reviewer', '?')} ({sel.get('reason', 'no reason given')})")]
    for f in data["flags"]:
        rows.append(("Flag", f))
    for c in data["approver_comments"]:
        rows.append(("Approver comment", c))
    rows.append(("Revision", str(data["revision"])))
    return rows


def render(data: dict, out_dir: Path | None = None) -> dict[str, str]:
    out_dir = (out_dir or config.REPORT_DIR) / data["job_id"]
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"report_rev{data['revision']}"
    rows = _lines(data)
    paths = {"pdf": out_dir / f"{stem}.pdf", "xlsx": out_dir / f"{stem}.xlsx", "docx": out_dir / f"{stem}.docx"}

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
    styles = getSampleStyleSheet()
    story = [Paragraph("Quality Inspection Report", styles["Title"]), Spacer(1, 12),
             Table([[k, Paragraph(v.replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"])]
                    for k, v in rows], colWidths=[110, 380])]
    SimpleDocTemplate(str(paths["pdf"]), pagesize=A4).build(story)

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Field", "Value"])
    for r in rows:
        ws.append(list(r))
    wb.save(paths["xlsx"])

    import docx
    doc = docx.Document()
    doc.add_heading("Quality Inspection Report", 1)
    table = doc.add_table(rows=0, cols=2)
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v
    doc.save(str(paths["docx"]))
    return {k: str(v) for k, v in paths.items()}
