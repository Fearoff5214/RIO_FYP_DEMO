"""Report rendering: PDF (ReportLab), Excel (openpyxl), Word (python-docx)."""
from __future__ import annotations

from pathlib import Path
from datetime import datetime, timezone


REPORTS_DIR = Path("tmp/reports")


async def render_report(report: dict) -> dict[str, str]:
    """
    Render report in all three formats.
    Returns {pdf, xlsx, docx} URI dict.
    """
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    job_id = report.get("job_id", "unknown")

    pdf_path = await _render_pdf(report, job_id)
    xlsx_path = await _render_xlsx(report, job_id)
    docx_path = await _render_docx(report, job_id)

    return {
        "pdf": str(pdf_path),
        "xlsx": str(xlsx_path),
        "docx": str(docx_path),
    }


async def _render_pdf(report: dict, job_id: str) -> Path:
    """Render PDF report using ReportLab."""
    out = REPORTS_DIR / f"{job_id}_report_v{report.get('version',1)}.pdf"

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        )

        doc = SimpleDocTemplate(str(out), pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        # Header
        story.append(Paragraph("NOVALISEO — Quality Inspection Report", styles["Title"]))
        story.append(Paragraph("Intelligent Inspection. Accurate Insights. Superior Quality.", styles["Normal"]))
        story.append(Spacer(1, 12))

        # Part details
        part = report.get("part", {}) or {}
        story.append(Paragraph("Part Details", styles["Heading2"]))
        part_data = [
            ["Part Number", part.get("part_number", "")],
            ["Part Name", part.get("part_name", "")],
            ["Drawing Rev", part.get("drawing_rev", "")],
            ["Material", part.get("material", "")],
            ["Process", part.get("process", "")],
            ["Report Date", report.get("generated_at", "")[:10]],
            ["Inspector", report.get("inspector", "System AI + Human")],
        ]
        t = Table(part_data, colWidths=[120, 300])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t)
        story.append(Spacer(1, 12))

        # Inspection summary
        story.append(Paragraph("Inspection Summary", styles["Heading2"]))
        summary = report.get("inspection_summary", {})
        status_data = [
            ["Check", "Result"],
            ["Dimensional Inspection", summary.get("dimensional", "N/A")],
            ["Defect Detection", summary.get("defect_detection", "N/A")],
        ]
        t2 = Table(status_data, colWidths=[200, 200])
        t2.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t2)
        story.append(Spacer(1, 12))

        # Defects
        defects = report.get("defects", []) or []
        if defects:
            story.append(Paragraph("Detected Defects", styles["Heading2"]))
            d_data = [["Type", "Severity", "Size (mm²)", "Confidence"]]
            for d in defects:
                d_data.append([
                    d.get("defect_type", ""),
                    d.get("severity", ""),
                    str(round(d.get("size_mm2", 0), 1)),
                    f"{d.get('confidence', 0):.0%}",
                ])
            t3 = Table(d_data, colWidths=[100, 100, 100, 100])
            t3.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.darkred),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(t3)
            story.append(Spacer(1, 12))

        # Root cause
        rca = report.get("root_cause") or {}
        if rca:
            story.append(Paragraph("Root Cause Analysis", styles["Heading2"]))
            story.append(Paragraph(f"<b>Root Cause:</b> {rca.get('cause', '')}", styles["Normal"]))
            story.append(Paragraph(f"<b>Effect/Impact:</b> {rca.get('effect', '')}", styles["Normal"]))
            story.append(Spacer(1, 6))

            corrective = rca.get("corrective_actions", [])
            if corrective:
                story.append(Paragraph("Corrective Actions:", styles["Heading3"]))
                for a in corrective:
                    story.append(Paragraph(f"• {a}", styles["Normal"]))

            preventive = rca.get("preventive_actions", [])
            if preventive:
                story.append(Paragraph("Preventive Actions:", styles["Heading3"]))
                for a in preventive:
                    story.append(Paragraph(f"• {a}", styles["Normal"]))

        # Footer
        story.append(Spacer(1, 24))
        story.append(Paragraph("— NOVALISEO · AI + Human Approved —", styles["Normal"]))

        doc.build(story)

    except ImportError:
        # Fallback: write a plain text file
        out.with_suffix(".txt").write_text(f"Report for job {job_id}\n{str(report)}")
        out = out.with_suffix(".txt")

    return out


async def _render_xlsx(report: dict, job_id: str) -> Path:
    """Render Excel report with multiple sheets."""
    out = REPORTS_DIR / f"{job_id}_report_v{report.get('version',1)}.xlsx"

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()

        # Summary sheet
        ws_sum = wb.active
        ws_sum.title = "Summary"
        part = report.get("part", {}) or {}
        summary = report.get("inspection_summary", {})

        ws_sum.append(["NOVALISEO Quality Inspection Report"])
        ws_sum.append([])
        ws_sum.append(["Part Number", part.get("part_number", "")])
        ws_sum.append(["Part Name", part.get("part_name", "")])
        ws_sum.append(["Date", report.get("generated_at", "")[:10]])
        ws_sum.append(["Dimensional", summary.get("dimensional", "")])
        ws_sum.append(["Defect Detection", summary.get("defect_detection", "")])

        # Dimensions sheet
        ws_dim = wb.create_sheet("Dimensions")
        ws_dim.append(["Feature ID", "Type", "Nominal", "Measured", "Deviation", "Within Tolerance"])
        for m in (report.get("measurements") or []):
            ws_dim.append([
                m.get("feature_id"),
                m.get("feature_type"),
                m.get("nominal"),
                m.get("measured"),
                m.get("deviation"),
                str(m.get("within_tolerance")),
            ])

        # Defects sheet
        ws_def = wb.create_sheet("Defects")
        ws_def.append(["Type", "Severity", "Size mm2", "Confidence", "Source"])
        for d in (report.get("defects") or []):
            ws_def.append([
                d.get("defect_type"),
                d.get("severity"),
                d.get("size_mm2"),
                d.get("confidence"),
                d.get("source"),
            ])

        # RCA sheet
        rca = report.get("root_cause") or {}
        if rca:
            ws_rca = wb.create_sheet("RCA")
            ws_rca.append(["Root Cause", rca.get("cause", "")])
            ws_rca.append(["Effect", rca.get("effect", "")])
            ws_rca.append(["Corrective Actions"])
            for a in rca.get("corrective_actions", []):
                ws_rca.append(["", a])
            ws_rca.append(["Preventive Actions"])
            for a in rca.get("preventive_actions", []):
                ws_rca.append(["", a])

        wb.save(str(out))

    except ImportError:
        out.write_text(f"Excel report for job {job_id}")

    return out


async def _render_docx(report: dict, job_id: str) -> Path:
    """Render Word report using python-docx."""
    out = REPORTS_DIR / f"{job_id}_report_v{report.get('version',1)}.docx"

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        doc.add_heading("NOVALISEO — Quality Inspection Report", 0)
        doc.add_paragraph("Intelligent Inspection. Accurate Insights. Superior Quality.")

        part = report.get("part", {}) or {}
        doc.add_heading("Part Details", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Value"
        for k, v in [
            ("Part Number", part.get("part_number", "")),
            ("Part Name", part.get("part_name", "")),
            ("Material", part.get("material", "")),
            ("Process", part.get("process", "")),
        ]:
            row = table.add_row().cells
            row[0].text = k
            row[1].text = str(v)

        summary = report.get("inspection_summary", {})
        doc.add_heading("Inspection Summary", level=1)
        doc.add_paragraph(f"Dimensional: {summary.get('dimensional', 'N/A')}")
        doc.add_paragraph(f"Defect Detection: {summary.get('defect_detection', 'N/A')}")

        rca = report.get("root_cause") or {}
        if rca:
            doc.add_heading("Root Cause Analysis", level=1)
            doc.add_paragraph(f"Root Cause: {rca.get('cause', '')}")
            doc.add_paragraph(f"Effect: {rca.get('effect', '')}")

        doc.add_heading("Approval", level=1)
        doc.add_paragraph("Approved By: _______________________")
        doc.add_paragraph("Date: _______________________")
        doc.add_paragraph("Signature: _______________________")

        doc.save(str(out))

    except ImportError:
        out.write_text(f"Word report for job {job_id}")

    return out
